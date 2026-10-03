import csv, hashlib, hmac, io, secrets
from datetime import timedelta
from cryptography.fernet import Fernet
from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.contrib.auth.password_validation import validate_password
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook, Workbook
from rest_framework.exceptions import ValidationError, PermissionDenied
from accounts.models import User
from institutions.models import StaffAssignment, TeachingAssignment, StudentMembership
from common.access import require_class, require_institution
from common.services import audit
from .models import Invitation, Prospect, ImportBatch


def cipher():
    if not settings.INVITATION_ENCRYPTION_KEY:
        raise ValidationError("INVITATION_ENCRYPTION_KEY не настроен")
    return Fernet(settings.INVITATION_ENCRYPTION_KEY.encode())


def digest(value):
    return hashlib.sha256(value.encode()).hexdigest()


def code_digest(invite, code):
    return hmac.new(
        settings.SECRET_KEY.encode(),
        (invite.token_hash + ":" + code).encode(),
        hashlib.sha256,
    ).hexdigest()


def live(invite):
    if invite.issued_by_id:
        if invite.role == "student":
            require_class(invite.issued_by, invite.classroom, True)
        else:
            require_institution(invite.issued_by, invite.institution, ["admin"])
    if invite.status != "pending":
        raise ValidationError("Приглашение недействительно")
    if invite.expires_at <= timezone.now():
        invite.status = "expired"
        invite.encrypted_token = ""
        invite.code_hash = ""
        invite.save()
        raise ValidationError("Срок приглашения истёк")


def invitation_url(invite):
    # Fragment does not appear in HTTP URLs, server access logs or Referer.
    if invite.status != "pending" or invite.expires_at <= timezone.now():
        return ""
    return (
        settings.WEB_URL
        + "/invite#"
        + cipher().decrypt(invite.encrypted_token.encode()).decode()
    )


def queue_mail(invite, encrypted_code=""):
    from common.outbox import enqueue

    suffix = digest(encrypted_code) if encrypted_code else "link"
    enqueue(
        "invitation",
        {"pk": str(invite.pk), "encrypted_code": encrypted_code},
        f"invite:{invite.pk}:{suffix}",
    )


@transaction.atomic
def issue(actor, institution, email, role, classroom=None, prospect=None):
    email = email.strip().lower()
    validate_email(email)
    if role == "student":
        if not classroom or not prospect:
            raise ValidationError(
                "Ученик должен быть связан с импортированной записью класса"
            )
        require_class(actor, classroom, True)
    else:
        if role not in ("teacher", "curator", "admin", "organizer", "workshop_teacher"):
            raise ValidationError("Недопустимая роль")
        require_institution(actor, institution, ["admin"])
        if role == "admin" and not actor.is_superuser:
            raise PermissionDenied("Администратора учреждения назначает платформа")
        if (
            role in ("organizer", "workshop_teacher")
            and institution.kind != "university"
        ):
            raise ValidationError("Роль мастер-класса назначается вузу")
    if classroom and classroom.institution_id != institution.pk:
        raise ValidationError("Класс другого учреждения")
    token = secrets.token_urlsafe(32)
    invite = Invitation.objects.create(
        email=email,
        issued_by=actor,
        institution=institution,
        classroom=classroom,
        prospect=prospect,
        role=role,
        token_hash=digest(token),
        encrypted_token=cipher().encrypt(token.encode()).decode(),
        expires_at=timezone.now() + timedelta(days=7),
    )
    audit(actor, "invitation.issued", invite, institution)
    queue_mail(invite)
    return invite


@transaction.atomic
def revoke(actor, invite):
    invite = Invitation.objects.select_for_update().get(pk=invite.pk)
    if invite.classroom_id:
        require_class(actor, invite.classroom, True)
    else:
        require_institution(actor, invite.institution, ["admin"])
    if invite.status == "accepted":
        raise ValidationError("Принятое приглашение нельзя отозвать")
    invite.status = "revoked"
    invite.encrypted_token = ""
    invite.code_hash = ""
    invite.save()
    audit(actor, "invitation.revoked", invite, invite.institution)
    return invite


@transaction.atomic
def correct_email(actor, prospect, new_email):
    prospect = Prospect.objects.select_for_update().get(pk=prospect.pk)
    require_class(actor, prospect.classroom, True)
    if prospect.user_id:
        raise ValidationError(
            "Почту зарегистрированного пользователя меняет владелец аккаунта"
        )
    new_email = new_email.strip().lower()
    validate_email(new_email)
    if (
        User.objects.filter(email__iexact=new_email).exists()
        or Prospect.objects.filter(email=new_email, classroom=prospect.classroom)
        .exclude(pk=prospect.pk)
        .exists()
    ):
        raise ValidationError("Почта уже используется; требуется разбор конфликта")
    old = prospect.email
    for invite in prospect.invitation_set.filter(status="pending"):
        revoke(actor, invite)
    prospect.email = new_email
    prospect.save()
    audit(
        actor,
        "prospect.email_changed",
        prospect,
        prospect.classroom.institution,
        {"old_email": old, "new_email": new_email},
    )
    return issue(
        actor,
        prospect.classroom.institution,
        new_email,
        "student",
        prospect.classroom,
        prospect,
    )


def request_code(token):
    with transaction.atomic():
        invite = (
            Invitation.objects.select_for_update()
            .filter(token_hash=digest(token))
            .first()
        )
        if not invite:
            raise ValidationError("Приглашение недействительно")
        live(invite)
        if invite.code_sent_at and invite.code_sent_at > timezone.now() - timedelta(
            minutes=1
        ):
            raise ValidationError("Повторите отправку через минуту")
        if invite.attempts >= 10:
            raise ValidationError("Слишком много попыток; обратитесь к преподавателю")
        code = str(secrets.randbelow(10**6)).zfill(6)
        invite.code_hash = code_digest(invite, code)
        invite.code_sent_at = timezone.now()
        invite.code_expires_at = timezone.now() + timedelta(minutes=10)
        invite.save()
        queue_mail(invite, cipher().encrypt(code.encode()).decode())


def accept(token, code, password, current_user=None):
    error = None
    user = None
    with transaction.atomic():
        invite = (
            Invitation.objects.select_for_update(of=("self",))
            .select_related("prospect", "classroom", "institution")
            .filter(token_hash=digest(token))
            .first()
        )
        if not invite:
            raise ValidationError("Приглашение недействительно")
        live(invite)
        if invite.attempts >= 10:
            raise ValidationError("Превышено число попыток")
        if (
            not invite.code_hash
            or not invite.code_expires_at
            or invite.code_expires_at <= timezone.now()
            or not hmac.compare_digest(invite.code_hash, code_digest(invite, code))
        ):
            invite.attempts += 1
            invite.save()
            error = "Неверный или просроченный код подтверждения"
        else:
            user = User.objects.filter(email__iexact=invite.email).first()
            if user:
                if (
                    not current_user
                    or not current_user.is_authenticated
                    or current_user.pk != user.pk
                ):
                    raise PermissionDenied(
                        "Войдите в существующий аккаунт с этой почтой"
                    )
                if not user.email_verified_at:
                    user.email_verified_at = timezone.now()
                    user.save(update_fields=["email_verified_at"])
            else:
                user = User(username=secrets.token_hex(16), email=invite.email)
                if invite.prospect:
                    user.first_name = invite.prospect.full_name[:150]
                    user.display_name = invite.prospect.full_name
                validate_password(password, user)
                user.set_password(password)
                user.email_verified_at = timezone.now()
                user.save()
            if invite.role == "student":
                StudentMembership.objects.get_or_create(
                    user=user,
                    classroom=invite.classroom,
                    defaults={"started_at": timezone.localdate()},
                )
                invite.prospect.user = user
                invite.prospect.save()
            else:
                staff, _ = StaffAssignment.objects.get_or_create(
                    user=user, institution=invite.institution, role=invite.role
                )
                if not staff.active:
                    staff.active = True
                    staff.save()
                if invite.role == "teacher" and invite.classroom:
                    assignment, _ = TeachingAssignment.objects.get_or_create(
                        staff=staff, classroom=invite.classroom
                    )
                    if not assignment.active:
                        assignment.active = True
                        assignment.save()
            invite.status = "accepted"
            invite.accepted_by = user
            invite.encrypted_token = ""
            invite.code_hash = ""
            invite.save()
            audit(user, "invitation.accepted", invite, invite.institution)
    if error:
        raise ValidationError(error)
    return user


def preview_import(actor, classroom, upload):
    require_class(actor, classroom, True)
    if upload.size > 2 * 1024 * 1024:
        raise ValidationError("Максимум 2 МБ")
    if upload.name.lower().endswith(".xlsx"):
        # Reject oversized expanded archives before parsing.
        import zipfile

        with zipfile.ZipFile(upload) as archive:
            if sum(i.file_size for i in archive.infolist()) > 10 * 1024 * 1024:
                raise ValidationError("Слишком большой XLSX")
        upload.seek(0)
        wb = load_workbook(upload, read_only=True, data_only=False)
        stream = wb.active.iter_rows(values_only=True)
    elif upload.name.lower().endswith(".csv"):
        stream = csv.reader(io.StringIO(upload.read().decode("utf-8-sig")))
    else:
        raise ValidationError("Требуется CSV или XLSX")
    header = next(stream, None)
    if not header or list(header)[:2] != ["ФИО", "email"]:
        raise ValidationError("Колонки должны называться ФИО и email")
    rows = []
    errors = []
    seen = set()
    for n, row in enumerate(stream, 2):
        if n > 501:
            raise ValidationError("Максимум 500 строк")
        name = str(row[0] or "").strip() if row else ""
        email = str(row[1] or "").strip().lower() if len(row) > 1 else ""
        problem = ""
        if not name or len(name) > 200:
            problem = "ФИО обязательно, максимум 200 символов"
        try:
            validate_email(email)
        except DjangoValidationError:
            problem = "Некорректный email"
        if email in seen:
            problem = "Дубликат email внутри файла"
        if User.objects.filter(email__iexact=email).exists():
            problem = "Email зарегистрирован: требуется отдельное подтверждённое присоединение"
        existing = Prospect.objects.filter(email=email, classroom=classroom).first()
        if existing and existing.full_name != name:
            problem = "Email связан с другим ФИО"
        if problem:
            errors.append({"row": n, "reason": problem})
        rows.append(
            {"row": n, "full_name": name, "email": email, "existing": bool(existing)}
        )
        seen.add(email)
    return ImportBatch.objects.create(
        actor=actor,
        classroom=classroom,
        rows=rows,
        errors=errors,
        expires_at=timezone.now() + timedelta(hours=1),
    )


@transaction.atomic
def apply_import(actor, batch):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    require_class(actor, batch.classroom, True)
    if batch.actor_id != actor.pk:
        raise PermissionDenied("Предпросмотр другого сотрудника")
    if batch.errors or batch.expires_at <= timezone.now():
        raise ValidationError("Предпросмотр содержит ошибки или истёк")
    if batch.applied_at:
        return {"created": 0, "skipped": len(batch.rows)}
    created = 0
    skipped = 0
    for row in batch.rows:
        if User.objects.filter(email__iexact=row["email"]).exists():
            raise ValidationError("Email уже зарегистрирован; повторите предпросмотр")
        p, new = Prospect.objects.get_or_create(
            classroom=batch.classroom,
            email=row["email"],
            defaults={"full_name": row["full_name"]},
        )
        if p.full_name != row["full_name"]:
            raise ValidationError("Конфликт ФИО; повторите предпросмотр")
        if new:
            issue(
                actor,
                batch.classroom.institution,
                p.email,
                "student",
                batch.classroom,
                p,
            )
            created += 1
        else:
            skipped += 1
    batch.applied_at = timezone.now()
    batch.save()
    audit(
        actor,
        "class.imported",
        batch,
        batch.classroom.institution,
        {"created": created, "skipped": skipped},
    )
    return {"created": created, "skipped": skipped}


def safe_cell(value):
    return (
        "'" + value
        if isinstance(value, str)
        and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
        else value
    )


def export_class(actor, classroom):
    require_class(actor, classroom, True)
    expire()
    wb = Workbook()
    ws = wb.active
    ws.title = "Приглашения"
    ws.append(
        ["ФИО", "Email", "Ссылка приглашения", "Статус приглашения", "Срок действия"]
    )
    for p in Prospect.objects.filter(classroom=classroom).order_by("full_name"):
        inv = p.invitation_set.order_by("-created_at").first()
        ws.append(
            [
                safe_cell(p.full_name),
                safe_cell(p.email),
                invitation_url(inv) if inv and not p.user_id else "",
                inv.status if inv else "",
                inv.expires_at.isoformat() if inv else "",
            ]
        )
    for m in classroom.studentmembership_set.select_related("user"):
        if not Prospect.objects.filter(classroom=classroom, user=m.user).exists():
            ws.append(
                [
                    safe_cell(m.user.get_full_name()),
                    safe_cell(m.user.email),
                    "",
                    "registered",
                    "",
                ]
            )
    output = io.BytesIO()
    wb.save(output)
    audit(actor, "class.invitation_exported", classroom, classroom.institution)
    return output.getvalue()


def expire():
    return Invitation.objects.filter(
        status="pending", expires_at__lte=timezone.now()
    ).update(status="expired", encrypted_token="", code_hash="")
