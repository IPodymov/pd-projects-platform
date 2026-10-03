from celery import shared_task
from common.access import project_access
from rest_framework.exceptions import PermissionDenied
from .models import DocumentVersion, Comparison


@shared_task(time_limit=35, soft_time_limit=30)
def extract_text(pk):
    v = DocumentVersion.objects.get(pk=pk)
    if v.extraction_status == "done":
        return
    try:
        with v.file.open("rb") as f:
            if v.filename.lower().endswith(".docx"):
                from docx import Document
                from docx.text.paragraph import Paragraph
                from docx.table import Table

                doc = Document(f)
                blocks = []
                for child in doc.element.body:
                    tag = child.tag.rsplit("}", 1)[-1]
                    if tag == "p":
                        blocks.append(
                            {"kind": "paragraph", "text": Paragraph(child, doc).text}
                        )
                    if tag == "tbl":
                        rows = [
                            [c.text for c in row.cells]
                            for row in Table(child, doc).rows
                        ]
                        blocks.append(
                            {
                                "kind": "table",
                                "rows": rows,
                                "text": "\n".join(" | ".join(row) for row in rows),
                            }
                        )
            else:
                from pptx import Presentation

                blocks = []
                for n, slide in enumerate(Presentation(f).slides, 1):
                    text = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            text.append(shape.text)
                        if shape.has_table:
                            text.extend(
                                " | ".join(cell.text for cell in row.cells)
                                for row in shape.table.rows
                            )
                    blocks.append(
                        {
                            "kind": "slide",
                            "index": n,
                            "blocks": text,
                            "text": "\n".join(text),
                        }
                    )
        text = "\n".join(b["text"] for b in blocks)
        if len(text) > 2_000_000:
            raise ValueError("Extraction size limit")
        DocumentVersion.objects.filter(pk=pk).update(
            text_content=text, structure=blocks, extraction_status="done"
        )
    except Exception:
        DocumentVersion.objects.filter(pk=pk).update(extraction_status="failed")
        raise


@shared_task(time_limit=35, soft_time_limit=30)
def run_comparison(pk):
    from .services import compare

    obj = Comparison.objects.select_related(
        "old__document__project", "new__document__project", "initiator"
    ).get(pk=pk)
    if obj.status in ("succeeded", "failed"):
        return
    try:
        project_access(obj.initiator, obj.old.document.project)
        project_access(obj.initiator, obj.new.document.project)
    except PermissionDenied:
        Comparison.objects.filter(pk=pk).update(
            status="failed", error_code="access_revoked"
        )
        return
    if any(v.extraction_status == "failed" for v in (obj.old, obj.new)):
        Comparison.objects.filter(pk=pk).update(
            status="failed", error_code="extraction_failed"
        )
        return
    if any(v.extraction_status != "done" for v in (obj.old, obj.new)):
        raise RuntimeError("Extraction is pending")
    Comparison.objects.filter(pk=pk).update(status="processing")
    try:
        result = compare(obj.initiator, obj.old, obj.new)
        result["old_structure"] = obj.old.structure
        result["new_structure"] = obj.new.structure
        Comparison.objects.filter(pk=pk).update(
            status="succeeded", result=result, error_code=""
        )
    except Exception:
        Comparison.objects.filter(pk=pk).update(
            status="failed", error_code="comparison_failed"
        )
        raise
