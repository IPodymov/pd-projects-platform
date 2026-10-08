from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.exceptions import PermissionDenied
from common.testing import PlatformCase
from projects.models import Project, ProjectMember
from .models import Document
from .services import upload_version, restore, compare


class DocumentTests(PlatformCase):
    def test_upload_immutable_history_restore_and_diff(self):
        project = Project.objects.create(title="Проект", classroom=self.classroom)
        ProjectMember.objects.create(project=project, user=self.student)
        doc = Document.objects.create(title="Документ", project=project)
        a = upload_version(
            self.student, doc, SimpleUploadedFile("a.txt", b"one\ntwo\n")
        )
        b = upload_version(
            self.student, doc, SimpleUploadedFile("a.txt", b"one\nthree\n")
        )
        self.assertEqual(b.previous_id, a.pk)
        self.assertNotEqual(a.checksum, b.checksum)
        diff = compare(self.student, a, b)
        self.assertTrue(
            any(
                line["kind"] == "add" and line["text"] == "+three"
                for line in diff["lines"]
            )
        )
        self.assertTrue(any(line.get("ranges") for line in diff["lines"]))
        c = restore(self.student, a)
        self.assertNotEqual(c.pk, a.pk)
        self.assertEqual(c.previous_id, b.pk)
        self.assertEqual(c.checksum, a.checksum)
        with self.assertRaises(ValueError):
            a.save()
        with self.assertRaises(PermissionDenied):
            upload_version(self.stranger, doc, SimpleUploadedFile("a.txt", b"hacked"))
        self.client.force_authenticate(self.stranger)
        self.assertEqual(
            self.client.get(
                "/api/document-versions/" + str(a.pk) + "/download/"
            ).status_code,
            404,
        )

    def test_background_comparison_idempotence_and_office_structure(self):
        from io import BytesIO
        from docx import Document as WordDocument
        from .services import request_comparison
        from .tasks import extract_text, run_comparison

        project = Project.objects.create(title="Документы", classroom=self.classroom)
        ProjectMember.objects.create(project=project, user=self.student)
        doc = Document.objects.create(title="Отчёт", project=project)

        def file(value):
            word = WordDocument()
            word.add_paragraph(value)
            table = word.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Параметр"
            table.cell(0, 1).text = value
            stream = BytesIO()
            word.save(stream)
            return SimpleUploadedFile("report.docx", stream.getvalue())

        old = upload_version(self.student, doc, file("10"))
        new = upload_version(self.student, doc, file("20"))
        extract_text(str(old.pk))
        extract_text(str(new.pk))
        old.refresh_from_db()
        new.refresh_from_db()
        job = request_comparison(self.student, old, new)
        self.assertEqual(request_comparison(self.student, old, new).pk, job.pk)
        run_comparison(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, "succeeded")
        self.assertEqual(job.result["old_structure"][1]["kind"], "table")
        self.assertTrue(any(row["kind"] == "add" for row in job.result["lines"]))
        with self.assertRaises(PermissionDenied):
            request_comparison(self.stranger, old, new)

    def test_pdf_upload_restore_and_no_text_comparison(self):
        from rest_framework.exceptions import ValidationError

        project = Project.objects.create(title="Материалы", classroom=self.classroom)
        ProjectMember.objects.create(project=project, user=self.student)
        document = Document.objects.create(title="PDF", project=project)
        with self.assertRaises(ValidationError):
            upload_version(
                self.student, document, SimpleUploadedFile("fake.pdf", b"html")
            )
        version = upload_version(
            self.student,
            document,
            SimpleUploadedFile("material.pdf", b"%PDF-1.4\n%%EOF"),
        )
        self.assertEqual(version.extraction_status, "unavailable")
        restored = restore(self.student, version)
        self.assertEqual(restored.checksum, version.checksum)
        with self.assertRaises(ValidationError):
            compare(self.student, version, restored)
        self.client.force_authenticate(self.student)
        response = self.client.get(f"/api/v1/document-versions/{version.pk}/download/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(response.streaming_content), b"%PDF-1.4\n%%EOF")
