"""Explicit development-only end-to-end storage/queue/mail check."""

import io, secrets, time, logging
from django.core.management.base import BaseCommand, CommandError
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.storage import default_storage
from docx import Document as Docx
import httpx
from accounts.models import User
from projects.models import Project
from documents.models import Document, Comparison
from documents.services import upload_version, compare, restore, request_comparison
from invitations.models import Prospect
from invitations.services import issue, revoke


class Command(BaseCommand):
    help = "Explicit dev smoke: actual S3, Celery extraction and Mailpit SMTP; run after demo_data"

    def handle(self, *args, **options):
        logging.getLogger("httpx").setLevel(logging.WARNING)
        if not settings.DEBUG:
            raise CommandError("Development only")
        if not default_storage.__class__.__module__.startswith("storages.backends.s3"):
            raise CommandError("Actual S3 storage required")
        actor = User.objects.filter(username="demo-teacher").first()
        project = Project.objects.filter(title="Умная теплица · DEMO").first()
        if not actor or not project:
            raise CommandError("Run demo_data explicitly first")
        doc = Document.objects.create(project=project, title="SMOKE temporary document")
        versions = []
        prospect = None
        invite = None
        try:
            a = upload_version(
                actor, doc, SimpleUploadedFile("smoke.txt", b"first\nline\n")
            )
            versions.append(a)
            b = upload_version(
                actor, doc, SimpleUploadedFile("smoke.txt", b"first\nchanged\n")
            )
            versions.append(b)
            assert any(x["kind"] == "add" for x in compare(actor, a, b)["lines"])
            with a.file.open("rb") as f:
                assert f.read() == b"first\nline\n"
            c = restore(actor, a)
            versions.append(c)
            assert c.checksum == a.checksum
            out = io.BytesIO()
            office = Docx()
            office.add_paragraph("Actual Celery extraction")
            office.save(out)
            d = upload_version(
                actor, doc, SimpleUploadedFile("smoke.docx", out.getvalue())
            )
            versions.append(d)
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                d.refresh_from_db()
                if d.extraction_status != "pending":
                    break
                time.sleep(0.25)
            assert (
                d.extraction_status == "done"
                and "Actual Celery extraction" in d.text_content
            )
            from pptx import Presentation

            ppt = Presentation()
            slide = ppt.slides.add_slide(ppt.slide_layouts[1])
            slide.shapes.title.text = "Actual PPTX extraction"
            stream = io.BytesIO()
            ppt.save(stream)
            e = upload_version(
                actor, doc, SimpleUploadedFile("smoke.pptx", stream.getvalue())
            )
            versions.append(e)
            comparison = request_comparison(actor, a, b)
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                e.refresh_from_db()
                comparison.refresh_from_db()
                if e.extraction_status == "done" and comparison.status == "succeeded":
                    break
                time.sleep(0.25)
            assert e.extraction_status == "done" and e.structure[0]["kind"] == "slide"
            assert comparison.status == "succeeded" and comparison.result["lines"]
            email = "smoke-" + secrets.token_hex(6) + "@example.test"
            prospect = Prospect.objects.create(
                full_name="SMOKE temporary recipient",
                email=email,
                classroom=project.classroom,
            )
            invite = issue(
                actor,
                project.classroom.institution,
                email,
                "student",
                project.classroom,
                prospect,
            )
            sent = False
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                response = httpx.get("http://mailpit:8025/api/v1/messages", timeout=5)
                response.raise_for_status()
                sent = any(
                    any(to.get("Address") == email for to in item.get("To", []))
                    for item in response.json().get("messages", [])
                )
                if sent:
                    break
                time.sleep(0.25)
            assert sent, "Mailpit has not received the queued invitation"
            self.stdout.write(
                "PASS: real S3 upload/download, background diff, restore, DOCX/PPTX Celery extraction, invitation SMTP to Mailpit"
            )
        finally:
            Comparison.objects.filter(old__document=doc).delete()
            for version in reversed(versions):
                version.file.delete(save=False)
                version.delete()
            doc.delete()
            if invite:
                revoke(actor, invite)
                invite.delete()
            if prospect:
                prospect.delete()
